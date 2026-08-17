"""Excel Parser Service for reading Microsoft Forms responses from Excel files.

This service reads data from Excel files synced from Microsoft Forms without
modifying the original file.
"""

import logging
from typing import List, Dict, Any, Optional
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class ExcelParserService:
    """Service for parsing Excel files containing Microsoft Forms responses."""
    
    def __init__(self, file_path: str, sheet_name: str = 'Responses', header_row: int = 1):
        """Initialize the Excel parser.
        
        Args:
            file_path: Path to the Excel file
            sheet_name: Name of the sheet to read
            header_row: Row number containing headers (1-indexed)
        """
        self.file_path = file_path
        self.sheet_name = sheet_name
        self.header_row = header_row
        self._workbook = None
    
    def _load_workbook(self):
        """Load the workbook if not already loaded."""
        if self._workbook is None:
            try:
                self._workbook = load_workbook(self.file_path, read_only=True, data_only=True)
                logger.info(f"Loaded workbook: {self.file_path}")
            except FileNotFoundError:
                logger.error(f"Excel file not found: {self.file_path}")
                raise
            except Exception as e:
                logger.error(f"Error loading workbook: {e}")
                raise
    
    def get_headers(self) -> List[str]:
        """Get the headers from the specified row.
        
        Returns:
            List of header names
        """
        self._load_workbook()
        
        if self.sheet_name not in self._workbook.sheetnames:
            raise ValueError(f"Sheet '{self.sheet_name}' not found in workbook")
        
        sheet = self._workbook[self.sheet_name]
        headers = []
        
        for cell in sheet[self.header_row]:
            if cell.value is not None:
                headers.append(str(cell.value).strip())
            else:
                headers.append('')
        
        logger.debug(f"Found {len(headers)} headers")
        return headers
    
    def get_all_rows(self) -> List[Dict[str, Any]]:
        """Get all rows from the sheet as dictionaries.
        
        Returns:
            List of dictionaries, one per row
        """
        self._load_workbook()
        headers = self.get_headers()
        
        if self.sheet_name not in self._workbook.sheetnames:
            raise ValueError(f"Sheet '{self.sheet_name}' not found in workbook")
        
        sheet = self._workbook[self.sheet_name]
        rows = []
        
        # Start from row after header
        for row_num, row in enumerate(sheet.iter_rows(min_row=self.header_row + 1, values_only=True), start=self.header_row + 1):
            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue
            
            row_data = {}
            for i, value in enumerate(row):
                if i < len(headers):
                    header = headers[i]
                    if header:  # Only add if header exists
                        row_data[header] = value
            
            if row_data:  # Only add non-empty rows
                row_data['_row_number'] = row_num
                rows.append(row_data)
        
        logger.info(f"Retrieved {len(rows)} rows from sheet '{self.sheet_name}'")
        return rows
    
    def get_new_responses_since(self, last_row: int) -> List[Dict[str, Any]]:
        """Get rows that were added after a specific row number.
        
        Args:
            last_row: The last known row number
            
        Returns:
            List of new row dictionaries
        """
        all_rows = self.get_all_rows()
        return [row for row in all_rows if row.get('_row_number', 0) > last_row]
    
    def count_rows(self) -> int:
        """Count the total number of data rows (excluding header).
        
        Returns:
            Number of data rows
        """
        return len(self.get_all_rows())
    
    def close(self):
        """Close the workbook if open."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None
            logger.debug("Workbook closed")
    
    def __enter__(self):
        """Context manager entry."""
        self._load_workbook()
        return self
    
    def __exit__(self, exc_type, exc_val, exc_tb):
        """Context manager exit."""
        self.close()
